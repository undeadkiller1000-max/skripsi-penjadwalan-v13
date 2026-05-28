"""
ui/exporter.py
==============
Export hasil optimasi ke file Excel (.xlsx) dengan 4 sheet.
"""

import io
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from calendar_utils import menit_ke_tanggal_waktu, format_durasi
from config import STASIUN

WARNA_HEADER_BIRU  = "1F4E79"
WARNA_HEADER_HIJAU = "1E5631"
WARNA_HEADER_ABU   = "4D4D4D"
WARNA_TERLAMBAT_BG = "FDECEA"
WARNA_TEPAT_BG     = "EAF7EC"


def export_ke_excel(
    hasil_pemenang: list[dict],
    hasil_fcfs: list[dict],
    tanggal_mulai: date,
    nama_metode_pemenang: str = "Optimasi",
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _buat_sheet_laporan_manajemen(wb, hasil_pemenang, tanggal_mulai, nama_metode_pemenang)
    _buat_sheet_jadwal_stasiun(wb, hasil_pemenang, tanggal_mulai)
    _buat_sheet_detail_jadwal(wb, hasil_pemenang, tanggal_mulai, f"Detail {nama_metode_pemenang}", WARNA_HEADER_HIJAU)
    _buat_sheet_detail_jadwal(wb, hasil_fcfs, tanggal_mulai, "Detail FCFS", WARNA_HEADER_ABU)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _buat_sheet_laporan_manajemen(wb, hasil_list, tanggal_mulai, nama_metode):
    ws = wb.create_sheet("Laporan Manajemen")
    headers = ["ID Pesanan","Produk","Unit","Prioritas","Deadline (Tgl)",
               "Deadline (Mnt)","Selesai (Tgl)","Selesai (Jam)","Selesai (Mnt)",
               "Tardiness (Mnt)","Tardiness (Hari)","W.Tardiness","Status"]
    _tulis_header(ws, headers, WARNA_HEADER_BIRU)
    for i, h in enumerate(hasil_list, start=2):
        tgl_s, wkt_s = menit_ke_tanggal_waktu(tanggal_mulai, h["completion_time"])
        row = [
            h["id_pesanan"], h["jenis_produk"].capitalize(), h["jumlah_unit"],
            h["prioritas"], h["deadline_tgl"].strftime("%d/%m/%Y"), h["deadline_mnt"],
            tgl_s.strftime("%d/%m/%Y"), wkt_s, round(h["completion_time"],1),
            round(h["tardiness"],1), format_durasi(h["tardiness"]),
            round(h["weighted_tardiness"],1),
            "TERLAMBAT" if h["terlambat"] else "Tepat Waktu",
        ]
        ws.append(row)
        bg = WARNA_TERLAMBAT_BG if h["terlambat"] else WARNA_TEPAT_BG
        fill = PatternFill(fill_type="solid", fgColor=bg)
        for col in range(1, len(headers)+1):
            ws.cell(row=i, column=col).fill = fill
    n = len(hasil_list)+2
    ws.cell(row=n, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=n, column=10, value=round(sum(h["tardiness"] for h in hasil_list),1)).font = Font(bold=True)
    ws.cell(row=n, column=12, value=round(sum(h["weighted_tardiness"] for h in hasil_list),1)).font = Font(bold=True)
    n_terlambat = sum(1 for h in hasil_list if h["terlambat"])
    ws.cell(row=n, column=13, value=f"{n_terlambat} terlambat / {len(hasil_list)} job").font = Font(bold=True)
    _auto_width(ws)


def _buat_sheet_jadwal_stasiun(wb, hasil_list, tanggal_mulai):
    ws = wb.create_sheet("Jadwal per Stasiun")
    headers = ["Stasiun","Resource #","ID Pesanan","Produk","Unit",
               "Tgl Mulai","Jam Mulai","Tgl Selesai","Jam Selesai",
               "Durasi (Mnt)","Prioritas","Status"]
    _tulis_header(ws, headers, WARNA_HEADER_HIJAU)
    semua_ops = []
    for h in hasil_list:
        for st, ops in h["schedule"].items():
            tgl_m, wkt_m = menit_ke_tanggal_waktu(tanggal_mulai, ops["start"])
            tgl_s, wkt_s = menit_ke_tanggal_waktu(tanggal_mulai, ops["end"])
            semua_ops.append({
                "st":st,"resource":ops["resource"],"id_pesanan":h["id_pesanan"],
                "produk":h["jenis_produk"].capitalize(),"unit":h["jumlah_unit"],
                "tgl_m":tgl_m.strftime("%d/%m/%Y"),"wkt_m":wkt_m,
                "tgl_s":tgl_s.strftime("%d/%m/%Y"),"wkt_s":wkt_s,
                "durasi":round(ops["end"]-ops["start"],1),"prioritas":h["prioritas"],
                "terlambat":h["terlambat"],"start_mnt":ops["start"],
            })
    semua_ops.sort(key=lambda x:(x["st"],x["resource"],x["start_mnt"]))
    for row_idx, ops in enumerate(semua_ops, start=2):
        ws.append([
            f"St.{ops['st']} - {STASIUN[ops['st']]}",ops["resource"],
            ops["id_pesanan"],ops["produk"],ops["unit"],
            ops["tgl_m"],ops["wkt_m"],ops["tgl_s"],ops["wkt_s"],
            ops["durasi"],ops["prioritas"],
            "TERLAMBAT" if ops["terlambat"] else "Tepat Waktu",
        ])
        bg = WARNA_TERLAMBAT_BG if ops["terlambat"] else WARNA_TEPAT_BG
        fill = PatternFill(fill_type="solid", fgColor=bg)
        for col in range(1, len(headers)+1):
            ws.cell(row=row_idx, column=col).fill = fill
    _auto_width(ws)


def _buat_sheet_detail_jadwal(wb, hasil_list, tanggal_mulai, nama_sheet, warna_header):
    ws = wb.create_sheet(nama_sheet[:31])
    headers = ["ID Pesanan","Produk","Unit","Prioritas","Stasiun","Resource #",
               "Mulai (Tgl)","Mulai (Jam)","Selesai (Tgl)","Selesai (Jam)",
               "Durasi (Mnt)","Status"]
    _tulis_header(ws, headers, warna_header)
    row_idx = 2
    for h in hasil_list:
        for st in h["routing"]:
            ops = h["schedule"][st]
            tgl_m, wkt_m = menit_ke_tanggal_waktu(tanggal_mulai, ops["start"])
            tgl_s, wkt_s = menit_ke_tanggal_waktu(tanggal_mulai, ops["end"])
            ws.append([
                h["id_pesanan"],h["jenis_produk"].capitalize(),h["jumlah_unit"],
                h["prioritas"],f"St.{st} - {STASIUN[st]}",ops["resource"],
                tgl_m.strftime("%d/%m/%Y"),wkt_m,tgl_s.strftime("%d/%m/%Y"),wkt_s,
                round(ops["end"]-ops["start"],1),
                "TERLAMBAT" if h["terlambat"] else "Tepat Waktu",
            ])
            bg = WARNA_TERLAMBAT_BG if h["terlambat"] else WARNA_TEPAT_BG
            fill = PatternFill(fill_type="solid", fgColor=bg)
            for col in range(1, len(headers)+1):
                ws.cell(row=row_idx, column=col).fill = fill
            row_idx += 1
    _auto_width(ws)


def _tulis_header(ws, headers, warna_hex):
    fill  = PatternFill(fill_type="solid", fgColor=warna_hex)
    font  = Font(bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.append(headers)
    for col_idx in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = fill; cell.font = font; cell.alignment = align
    ws.row_dimensions[1].height = 30


def _auto_width(ws, min_w=10, max_w=40):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value)) if cell.value else 0)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len+2, min_w), max_w)
